from __future__ import annotations

import io
import re
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill


st.set_page_config(
    page_title="Noon → Zoho Import Converter",
    page_icon="📦",
    layout="wide",
)

REQUIRED_COLUMNS = {
    "Reference Nr",
    "Transaction Type",
    "Order Nr",
    "Item Nr",
    "Transaction Date",
    "Title",
    "SKUs",
    "Net Proceeds",
}

INVOICE_COLUMNS = [
    "Invoice Date", "Invoice Number", "PurchaseOrder", "item name", "SKU",
    "Quantity", "Item Price", "VAT Treatment", "Accounts Receivable",
    "Is Inclusive Tax", "Template Name", "Currency Code", "Exchange Rate",
    "Account", "Item Tax", "Item Tax %", "Item Tax Type", "Item Type",
    "Noon Order Nr", "Reference Nr", "Customer Name", "Location Name",
    "Line Item Location Name", "Status",
]

EXPENSE_REQUIRED_COLUMNS = {
    "Transaction Type",
    "Document Date",
    "Invoice Nr",
    "Source Doc Nr",
    "Description",
    "VAT Rate",
    "Price Including VAT (Document Currency)",
}

# Expense output columns follow your template.
# Invoice Nr has been added as requested.
EXPENSE_COLUMNS = [
    "Expense Date",
    "Invoice Nr",
    "Expense Account",
    "Paid Through",
    "Location Name",
    "VAT Treatment",
    "Place Of Supply",
    "Currency Code",
    "Exchange Rate",
    "Is Inclusive Tax",
    "Tax Name",
    "Tax Percentage",
    "Tax Type",
    "Expense Type",
    "Expense Amount",
    "CF.Reference No",
    "Source Doc Nr",
]

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


# ============================================================
# COMMON FUNCTIONS
# ============================================================

def read_upload(uploaded_file) -> pd.DataFrame:
    raw = uploaded_file.getvalue()

    if not raw or not raw.strip():
        raise ValueError(
            "The uploaded file is empty (0 bytes). "
            "Please re-export it from Noon and upload again."
        )

    suffix = uploaded_file.name.lower().rsplit(".", 1)[-1]
    is_real_xlsx = raw.startswith(b"PK")

    if suffix in {"xlsx", "xls"} or is_real_xlsx:
        return pd.read_excel(io.BytesIO(raw))

    if suffix in {"csv", "tsv", "txt"}:
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encodings = ("utf-16",)
        else:
            encodings = ("utf-8-sig", "utf-8", "cp1252")

        last_error = None

        for encoding in encodings:
            try:
                return pd.read_csv(
                    io.BytesIO(raw),
                    encoding=encoding,
                    sep=None,
                    engine="python",
                )
            except UnicodeDecodeError as error:
                last_error = error
            except pd.errors.EmptyDataError:
                raise ValueError(
                    "The file has no header row or data. "
                    "Please re-export it from Noon."
                )

        raise ValueError(
            f"Could not read the CSV encoding: {last_error}"
        )

    raise ValueError("Please upload a CSV, XLSX, or XLS file.")


def excel_bytes(
    data: pd.DataFrame,
    date_columns: list[str],
) -> bytes:
    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
        datetime_format="DD/MM/YYYY",
    ) as writer:

        data.to_excel(
            writer,
            index=False,
            sheet_name="Import",
        )

        sheet = writer.book["Import"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        header_fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )

        for cell in sheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center"
            )

        for cells in sheet.columns:
            values = [
                "" if cell.value is None else str(cell.value)
                for cell in cells
            ]

            sheet.column_dimensions[
                cells[0].column_letter
            ].width = min(
                max(max(map(len, values)) + 2, 12),
                45,
            )

        headers = {
            cell.value: cell.column
            for cell in sheet[1]
        }

        for name in date_columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(
                    row=row,
                    column=headers[name],
                ).number_format = "DD/MM/YYYY"

    return output.getvalue()


def safe_name(value: str) -> str:
    return re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        value,
    ).strip("._")


def reference_sort_key(reference: str):
    match = re.search(r"(\d{8})$", reference)

    if match:
        try:
            return datetime.strptime(
                match.group(1),
                "%Y%m%d",
            )
        except ValueError:
            pass

    return datetime.min


# ============================================================
# SALES & RETURNS FUNCTIONS
# Existing sales and returns processing is preserved.
# ============================================================

def clean_source(data: pd.DataFrame) -> pd.DataFrame:
    data = data.copy()
    data.columns = data.columns.astype(str).str.strip()

    missing = sorted(
        REQUIRED_COLUMNS - set(data.columns)
    )

    if missing:
        raise ValueError(
            "Missing required column(s): "
            + ", ".join(missing)
        )

    # Keep only Order and Order Update transaction types.
    transaction_type = (
        data["Transaction Type"]
        .astype("string")
        .str.strip()
        .str.casefold()
        .str.replace(
            r"[\s_\-]+",
            " ",
            regex=True,
        )
    )

    data = data[
        transaction_type.isin(
            ["order", "order update"]
        )
    ].copy()

    if data.empty:
        raise ValueError(
            "No rows were found with Transaction Type "
            "'Order' or 'Order Update'."
        )

    data["Reference Nr"] = (
        data["Reference Nr"]
        .astype("string")
        .str.strip()
    )

    data = data[
        data["Reference Nr"].str.startswith(
            "PS",
            na=False,
        )
    ].copy()

    if data.empty:
        raise ValueError(
            "No rows were found where Reference Nr starts with PS."
        )

    data["Net Proceeds"] = pd.to_numeric(
        data["Net Proceeds"],
        errors="coerce",
    )

    invalid_amounts = data["Net Proceeds"].isna()

    if invalid_amounts.any():
        rows = ", ".join(
            map(
                str,
                (
                    data.index[invalid_amounts] + 2
                ).tolist()[:20],
            )
        )

        raise ValueError(
            "Net Proceeds is not numerical on source row(s): "
            + rows
        )

    data["Transaction Date"] = pd.to_datetime(
        data["Transaction Date"],
        errors="coerce",
        dayfirst=True,
    )

    invalid_dates = data["Transaction Date"].isna()

    if invalid_dates.any():
        rows = ", ".join(
            map(
                str,
                (
                    data.index[invalid_dates] + 2
                ).tolist()[:20],
            )
        )

        raise ValueError(
            "Transaction Date is invalid on source row(s): "
            + rows
        )

    for column in (
        "Order Nr",
        "Item Nr",
        "Title",
        "SKUs",
    ):
        blank = (
            data[column].isna()
            | data[column].astype(str).str.strip().eq("")
        )

        if blank.any():
            rows = ", ".join(
                map(
                    str,
                    (
                        data.index[blank] + 2
                    ).tolist()[:20],
                )
            )

            raise ValueError(
                f"{column} is blank on source row(s): "
                + rows
            )

    return data


def build_invoice_rows(
    data: pd.DataFrame,
) -> pd.DataFrame:

    result = pd.DataFrame({
        "Invoice Date": (
            data["Transaction Date"].astype(str)
        ),
        "Invoice Number": data["Item Nr"].astype(str),
        "PurchaseOrder": data["Item Nr"].astype(str),
        "item name": data["Title"],
        "SKU": data["SKUs"].astype(str),
        "Quantity": 1,
        "Item Price": data["Net Proceeds"],
        "VAT Treatment": "vat_not_registered",
        "Accounts Receivable": "Accounts Receivable",
        "Is Inclusive Tax": True,
        "Template Name": "Standard Template",
        "Currency Code": "AED",
        "Exchange Rate": 1,
        "Account": "Sales",
        "Item Tax": "Standard Rate",
        "Item Tax %": 5,
        "Item Tax Type": "ItemAmount",
        "Item Type": "goods",
        "Noon Order Nr": data["Order Nr"].astype(str),
        "Reference Nr": data["Reference Nr"].astype(str),
        "Customer Name": "Noon Receivables",
        "Location Name": "AlphaWolf Inventory (Main)",
        "Line Item Location Name": "Noon Inventory",
        "Status": "Sent",
    })

    return result[INVOICE_COLUMNS]


# ============================================================
# EXPENSES FUNCTIONS
# ============================================================

def clean_expense_source(
    data: pd.DataFrame,
) -> pd.DataFrame:

    data = data.copy()
    data.columns = data.columns.astype(str).str.strip()

    missing = sorted(
        EXPENSE_REQUIRED_COLUMNS - set(data.columns)
    )

    if missing:
        raise ValueError(
            "Missing required column(s): "
            + ", ".join(missing)
        )

    # Keep only Service Fee and Statement Fee.
    # Customer transactions are excluded.
    transaction_type = (
        data["Transaction Type"]
        .astype("string")
        .str.strip()
        .str.casefold()
    )

    data = data[
        transaction_type.isin(
            ["service fee", "statement fee"]
        )
    ].copy()

    if data.empty:
        raise ValueError(
            "No rows were found with Transaction Type "
            "'Service Fee' or 'Statement Fee'."
        )

    # Validate document dates.
    data["Document Date"] = pd.to_datetime(
        data["Document Date"],
        errors="coerce",
        dayfirst=True,
    )

    invalid_dates = data["Document Date"].isna()

    if invalid_dates.any():
        rows = ", ".join(
            map(
                str,
                (
                    data.index[invalid_dates] + 2
                ).tolist()[:20],
            )
        )

        raise ValueError(
            "Document Date is invalid on source row(s): "
            + rows
        )

    # Validate VAT rates.
    data["VAT Rate"] = pd.to_numeric(
        data["VAT Rate"],
        errors="coerce",
    )

    invalid_vat = data["VAT Rate"].isna()

    if invalid_vat.any():
        rows = ", ".join(
            map(
                str,
                (
                    data.index[invalid_vat] + 2
                ).tolist()[:20],
            )
        )

        raise ValueError(
            "VAT Rate is not numerical on source row(s): "
            + rows
        )

    # Validate expense amounts.
    amount_column = (
        "Price Including VAT (Document Currency)"
    )

    data[amount_column] = pd.to_numeric(
        data[amount_column],
        errors="coerce",
    )

    invalid_amounts = data[amount_column].isna()

    if invalid_amounts.any():
        rows = ", ".join(
            map(
                str,
                (
                    data.index[invalid_amounts] + 2
                ).tolist()[:20],
            )
        )

        raise ValueError(
            f"{amount_column} is not numerical on source row(s): "
            + rows
        )

    # Only the VAT rates specified in your rules are accepted.
    unsupported_vat = ~data["VAT Rate"].isin([0, 5])

    if unsupported_vat.any():
        values = sorted(
            data.loc[
                unsupported_vat,
                "VAT Rate",
            ].dropna().unique().tolist()
        )

        raise ValueError(
            "VAT Rate contains value(s) other than 0 or 5: "
            f"{values}"
        )

    return data


def build_expense_rows(
    data: pd.DataFrame,
) -> pd.DataFrame:

    descriptions = (
        data["Description"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    expense_accounts = []
    reference_numbers = []

    for description in descriptions:

        # PS descriptions:
        # PS number goes to CF.Reference No.
        # Text after the colon goes to Expense Account.
        match = re.match(
            r"^(PS\S*)\s*:\s*(.*)$",
            description,
            flags=re.IGNORECASE,
        )

        if match:
            reference_numbers.append(
                match.group(1).strip()
            )

            expense_accounts.append(
                match.group(2).strip()
            )

        else:
            # Non-PS descriptions remain unchanged.
            reference_numbers.append("")
            expense_accounts.append(description)

    amount_column = (
        "Price Including VAT (Document Currency)"
    )

    amounts = data[amount_column].astype(float).copy()

    # Damaged Returns Fee must always be negative.
    damaged_returns = (
        pd.Series(
            expense_accounts,
            index=data.index,
        )
        .str.strip()
        .str.casefold()
        .isin([
            "damaged returns fee",
            "damage penalty compensation",
        ])
    )

    amounts.loc[damaged_returns] = (
        -amounts.loc[damaged_returns].abs()
    )
    # VAT mapping: numeric values, with no percent sign.
    tax_names = data["VAT Rate"].map({
        5: "Standard Rate",
        0: "Zero Rate",
    })

    tax_percentages = data["VAT Rate"].astype(int)

    result = pd.DataFrame({
        "Expense Date": (
            data["Document Date"].dt.strftime("%d/%m/%Y")
        ),

        "Invoice Nr": (
            data["Invoice Nr"]
            .fillna("")
            .astype(str)
            .replace("nan", "")
        ),

        "Expense Account": expense_accounts,

        # Defaults copied from your expense template.
        "Paid Through": "Noon Clearing Account",
        "Location Name": "AlphaWolf Inventory (Main)",
        "VAT Treatment": "vat_registered",
        "Place Of Supply": "DU",
        "Currency Code": "AED",
        "Exchange Rate": 1,
        "Is Inclusive Tax": True,

        "Tax Name": tax_names,
        "Tax Percentage": tax_percentages,
        "Tax Type": "ItemAmount",
        "Expense Type": "service",

        "Expense Amount": amounts,
        "CF.Reference No": reference_numbers,

        "Source Doc Nr": (
            data["Source Doc Nr"]
            .fillna("")
            .astype(str)
            .replace("nan", "")
        ),
    })

    return result[EXPENSE_COLUMNS]


# ============================================================
# APP INTERFACE
# ============================================================

st.title("Noon → Zoho Books Import Converter")

# Slider to switch between the two pages.
page = st.segmented_control(
    "Choose converter",
    options=["Sales & Returns", "Expenses"],
    default="Sales & Returns",
    selection_mode="single",
    key="converter_page",
)

if page is None:
    page = "Sales & Returns"

# ============================================================
# PAGE 1: SALES & RETURNS
# ============================================================

if page == "Sales & Returns":

    st.caption(
        "Prepare invoice and sales-return import workbooks "
        "from a Noon Transaction View export."
    )

    uploaded = st.file_uploader(
        "Upload Noon Transaction View",
        type=["csv", "xlsx", "xls"],
        key="sales_returns_upload",
    )

    if uploaded is None:
        st.info(
            "Upload a Noon Transaction View file to begin."
        )
        st.stop()

    try:
        source = clean_source(
            read_upload(uploaded)
        )

    except Exception as error:
        st.error(str(error))
        st.stop()

    references = sorted(
        source["Reference Nr"]
        .dropna()
        .astype(str)
        .unique(),
        key=reference_sort_key,
        reverse=True,
    )

    selected_reference = st.selectbox(
        "Reference number",
        ["All"] + references,
        index=0,
    )

    if selected_reference == "All":
        selected = source.copy()
        export_name = "ALL_PS"

    else:
        selected = source[
            source["Reference Nr"] == selected_reference
        ].copy()

        export_name = safe_name(selected_reference)

    # Existing logic: positive = sales, negative = returns.
    # Zero-value rows are excluded from both exports.
    positive = selected[
        selected["Net Proceeds"] > 0
    ].copy()

    negative = selected[
        selected["Net Proceeds"] < 0
    ].copy()

    zero_count = int(
        (selected["Net Proceeds"] == 0).sum()
    )

    sales_total = positive["Net Proceeds"].sum()
    returns_total = negative["Net Proceeds"].sum()

    actual_total = selected["Net Proceeds"].sum()

    count_cols = st.columns(3)

    count_cols[0].metric(
        "Sales rows",
        len(positive),
    )

    count_cols[1].metric(
        "Return rows",
        len(negative),
    )

    count_cols[2].metric(
        "Zero rows discarded",
        zero_count,
    )

    amount_cols = st.columns(3)

    amount_cols[0].metric(
        "Sales total (positive)",
        f"AED {sales_total:,.2f}",
    )

    amount_cols[1].metric(
        "Returns total (negative)",
        f"AED {returns_total:,.2f}",
    )

    amount_cols[2].metric(
        "Actual total net proceeds",
        f"AED {actual_total:,.2f}",
    )

    st.subheader("Download Zoho import files")

    button_cols = st.columns(2)

    with button_cols[0]:

        if positive.empty:
            st.button(
                "Download Sales file",
                disabled=True,
                use_container_width=True,
                help="No rows with positive Net Proceeds.",
            )

        else:
            st.download_button(
                "Download Sales file",
                excel_bytes(
                    build_invoice_rows(positive),
                    ["Invoice Date"],
                ),
                file_name=(
                    f"invoice_import_{export_name}_sales.xlsx"
                ),
                mime=XLSX_MIME,
                type="primary",
                use_container_width=True,
            )

    with button_cols[1]:

        if negative.empty:
            st.button(
                "Download Returns file",
                disabled=True,
                use_container_width=True,
                help="No rows with negative Net Proceeds.",
            )

        else:
            st.download_button(
                "Download Returns file",
                excel_bytes(
                    build_invoice_rows(negative),
                    ["Invoice Date"],
                ),
                file_name=(
                    f"invoice_import_{export_name}_sales_returns.xlsx"
                ),
                mime=XLSX_MIME,
                use_container_width=True,
            )

    with st.expander("Preview selected rows"):

        st.dataframe(
            selected[
                [
                    "Transaction Date",
                    "Reference Nr",
                    "Order Nr",
                    "Item Nr",
                    "Title",
                    "SKUs",
                    "Net Proceeds",
                ]
            ],
            use_container_width=True,
            hide_index=True,
        )


# ============================================================
# PAGE 2: EXPENSES
# ============================================================

else:

    st.caption(
        "Prepare Zoho Books expense import workbooks "
        "from Noon invoices and credit notes."
    )

    expense_upload = st.file_uploader(
        "Upload Noon invoices and credit notes file",
        type=["csv", "xlsx", "xls"],
        key="expenses_upload",
    )

    if expense_upload is None:
        st.info(
            "Upload the Noon invoices and credit notes file "
            "to prepare the expense import."
        )
        st.stop()

    try:
        expense_source = clean_expense_source(
            read_upload(expense_upload)
        )

        expense_rows = build_expense_rows(
            expense_source
        )

    except Exception as error:
        st.error(str(error))
        st.stop()

    total_expenses = expense_rows["Expense Amount"].sum()

    expense_metrics = st.columns(2)

    expense_metrics[0].metric(
        "Expense rows",
        len(expense_rows),
    )

    expense_metrics[1].metric(
        "Total expense amount",
        f"AED {total_expenses:,.2f}",
    )

    st.subheader("Download Zoho expense import file")

    st.download_button(
        "Download Expenses file",
        excel_bytes(
            expense_rows,
            ["Expense Date"],
        ),
        file_name="noon_expenses_import.xlsx",
        mime=XLSX_MIME,
        type="primary",
        use_container_width=True,
    )

    with st.expander("Preview expense rows"):

        st.dataframe(
            expense_rows,
            use_container_width=True,
            hide_index=True,
        )